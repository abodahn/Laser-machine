"""Report builders and Excel / PDF renderers.

A report is just {title, subtitle, sections:[{name, columns, rows}]}; two renderers
turn that into .xlsx or .pdf, so adding a report means adding one query function.
"""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import alerts, analytics, db

BRAND = "T&C Laser Intelligence & Command Center"


def _rng(days):
    end = dt.date.today()
    return (end - dt.timedelta(days=days - 1)).isoformat(), end.isoformat()


def _sec(name, rows, columns=None):
    if not rows:
        return {"name": name, "columns": columns or [], "rows": []}
    cols = columns or list(rows[0].keys())
    return {"name": name, "columns": cols, "rows": [[r.get(c) for c in cols] for r in rows]}


# ------------------------------------------------------------------ reports

def daily_production(date=None):
    date = date or dt.date.today().isoformat()
    ov = analytics.department_overview(date)
    t = ov["totals"]
    summary = [{"Metric": "Units produced", "Value": t["units"]},
               {"Metric": "Jobs completed", "Value": t["jobs"]},
               {"Metric": "Target", "Value": t["target"] or "-"},
               {"Metric": "Target achieved %", "Value": t["target_pct"] or "-"},
               {"Metric": "Department utilization %", "Value": t["utilization"]},
               {"Metric": "Laser duty %", "Value": t["laser_duty"]},
               {"Metric": "Downtime hours", "Value": t["down_hours"]},
               {"Metric": "Idle hours", "Value": t["idle_hours"]},
               {"Metric": "Alarms", "Value": t["alarms"]},
               {"Metric": "Machines online", "Value": f"{t['machines_online']}/{t['machines_total']}"}]
    per = [{"Machine": m["name"], "Status": m["status"], "Units": m["units"], "Jobs": m["jobs"],
            "Target": m["target"], "Target %": m["target_pct"], "Utilization %": m["utilization"],
            "Performance %": m["performance"], "OEE %": m["oee"], "Health": m["health"],
            "Down min": m["down_minutes"], "Idle min": m["idle_minutes"], "Alarms": m["alarms"]}
           for m in ov["machines"]]
    shifts = analytics.shift_performance(days=1)
    return {"title": "Daily Laser Production Report", "subtitle": date,
            "sections": [_sec("Summary", summary), _sec("By machine", per),
                         _sec("By shift", [{k: (round(v, 2) if isinstance(v, float) else v)
                                            for k, v in s.items()} for s in shifts])]}


def shift_report(days=7):
    start, end = _rng(days)
    rows = analytics.shift_performance(days=days)
    agg = {}
    for r in rows:
        a = agg.setdefault(r["shift"] or "-", {"Shift": r["shift"] or "-", "Units": 0, "Jobs": 0,
                                               "Busy h": 0.0, "Idle h": 0.0, "Down h": 0.0, "Alarms": 0})
        a["Units"] += r["units"] or 0
        a["Jobs"] += r["jobs"] or 0
        a["Busy h"] += r["busy_hours"] or 0
        a["Idle h"] += r["idle_hours"] or 0
        a["Down h"] += r["down_hours"] or 0
        a["Alarms"] += r["alarms"] or 0
    for a in agg.values():
        for k in ("Busy h", "Idle h", "Down h"):
            a[k] = round(a[k], 1)
    return {"title": "Shift Performance Report", "subtitle": f"{start} to {end}",
            "sections": [_sec("Shift totals", list(agg.values())),
                         _sec("Detail", [{k: (round(v, 2) if isinstance(v, float) else v)
                                          for k, v in r.items()} for r in rows])]}


def utilization_report(days=30):
    start, end = _rng(days)
    rows = analytics.ranking(days=days, metric="utilization")["rows"]
    out = [{"Machine": r["name"], "Units": r["units"], "Jobs": r["jobs"],
            "Utilization %": _r(r["utilization"]), "Performance %": _r(r["performance"]),
            "OEE %": _r(r["oee"]), "Units/hour": _r(r["uph"]), "Laser duty %": _r(r["laser_duty"]),
            "Down min": _r(r["down_minutes"]), "Idle min": _r(r["idle_minutes"]),
            "Health": _r(r["health"])} for r in rows]
    cap = analytics.capacity(days=days)
    return {"title": "Machine Utilization Report", "subtitle": f"{start} to {end}",
            "sections": [_sec("Utilization", out),
                         _sec("Department capacity", [{"Metric": k, "Value": v} for k, v in cap.items()])]}


def downtime_report(days=30, machine_id=None):
    start, end = _rng(days)
    a = analytics.downtime_analysis(days=days, machine_id=machine_id)
    return {"title": "Downtime Report", "subtitle": f"{start} to {end}",
            "sections": [_sec("By category", [{k: _r(v) for k, v in r.items()} for r in a["by_kind"]]),
                         _sec("By machine", [{k: _r(v) for k, v in r.items()} for r in a["by_machine"]]),
                         _sec("Longest events", [{k: _r(v) for k, v in r.items()} for r in a["longest"]])]}


def alarm_report(days=30, machine_id=None):
    start, end = _rng(days)
    a = analytics.alarm_analysis(days=days, machine_id=machine_id)
    return {"title": "Alarm Report", "subtitle": f"{start} to {end}",
            "sections": [_sec("Top alarms", a["top"]), _sec("By machine", a["by_machine"]),
                         _sec("Daily trend", a["daily"])]}


def health_report():
    today = dt.date.today().isoformat()
    rows = []
    for m in db.q("SELECT id,name,model,serial_number,location FROM machines WHERE enabled=1 ORDER BY id"):
        d = db.q1("SELECT * FROM machine_daily WHERE machine_id=? AND production_date=?", (m["id"], today))
        mr = analytics.maintenance_risk(m["id"])
        hf = mr["health"]
        rows.append({"Machine": m["name"], "Model": m["model"], "Serial": m["serial_number"],
                     "Health": _r(d["health_score"]) if d else None,
                     "Trend/day": hf["slope_per_day"], "In 7 days": hf["in_7_days"],
                     "Risk": mr["level"], "Risk score": mr["risk"],
                     "Critical alarms 14d": mr["critical_alarms_14d"],
                     "Head drift %": mr["head_drift_pct"], "Temp slope": mr["temp_slope"],
                     "Reasons": "; ".join(mr["reasons"])})
    return {"title": "Machine Health Report", "subtitle": today,
            "sections": [_sec("Health & maintenance risk", rows)]}


def weekly_report():
    return _period_report(7, "Weekly Performance Report")


def monthly_report():
    return _period_report(30, "Monthly Management Report")


def _period_report(days, title):
    start, end = _rng(days)
    rank = analytics.ranking(days=days)["rows"]
    cap = analytics.capacity(days=days)
    bn = analytics.bottleneck(days=days)
    al = analytics.alarm_analysis(days=days)
    daily = [dict(r) for r in db.q(
        """SELECT production_date "Date", SUM(units) "Units", SUM(jobs) "Jobs",
                  ROUND(AVG(utilization),1) "Utilization %", ROUND(AVG(oee),1) "OEE %",
                  ROUND(SUM(down_seconds+offline_seconds)/3600.0,1) "Down h",
                  SUM(alarm_count) "Alarms"
           FROM machine_daily WHERE production_date >= ? GROUP BY production_date ORDER BY production_date""",
        (start,))]
    perf = [{"Machine": r["name"], "Units": r["units"], "Utilization %": _r(r["utilization"]),
             "OEE %": _r(r["oee"]), "Health": _r(r["health"]), "Down min": _r(r["down_minutes"]),
             "Alarms": r["alarms"]} for r in rank]
    return {"title": title, "subtitle": f"{start} to {end}",
            "sections": [_sec("Daily totals", daily), _sec("Machine performance", perf),
                         _sec("Capacity", [{"Metric": k, "Value": v} for k, v in cap.items()]),
                         _sec("Bottlenecks (lost hours)",
                              [{"Machine": r["name"], "Lost hours": _r(r["lost_hours"]),
                                "Down hours": _r(r["down_hours"]), "Idle hours": _r(r["idle_hours"]),
                                "Lost units est.": r["lost_units_estimate"]} for r in bn["rows"][:15]]),
                         _sec("Top alarms", al["top"][:15])]}


def loss_report(date=None):
    la = analytics.loss_analysis(date)
    return {"title": "Production Loss Analysis", "subtitle": la["date"],
            "sections": [_sec("Loss by cause", [{"Cause": k, "Lost units": v}
                                                for k, v in la["buckets"].items()]),
                         _sec("By machine", la["by_machine"])]}


def capacity_report(days=30):
    start, end = _rng(days)
    cap = analytics.capacity(days=days)
    bn = analytics.bottleneck(days=days)
    return {"title": "Capacity Analysis", "subtitle": f"{start} to {end}",
            "sections": [_sec("Capacity", [{"Metric": k, "Value": v} for k, v in cap.items()]),
                         _sec("Capacity loss by machine",
                              [{"Machine": r["name"], "Lost hours": _r(r["lost_hours"]),
                                "Idle hours": _r(r["idle_hours"]), "Down hours": _r(r["down_hours"]),
                                "Utilization %": _r(r["utilization"]),
                                "Lost units est.": r["lost_units_estimate"]} for r in bn["rows"]])]}


def maintenance_report(days=30):
    start, end = _rng(days)
    rows = []
    for m in db.q("SELECT id,name FROM machines WHERE enabled=1 ORDER BY id"):
        mr = analytics.maintenance_risk(m["id"])
        rows.append({"Machine": m["name"], "Risk": mr["level"], "Score": mr["risk"],
                     "Critical alarms 14d": mr["critical_alarms_14d"],
                     "Head drift %": mr["head_drift_pct"], "Temp slope degC/day": mr["temp_slope"],
                     "Health now": mr["health"]["current"], "Health in 7d": mr["health"]["in_7_days"],
                     "Findings": "; ".join(mr["reasons"]) or "-"})
    al = analytics.alarm_analysis(days=days)
    hcm = [dict(r) for r in db.q(
        """SELECT m.name "Machine", h.ts "Date", ROUND(h.slope_x,4) "Slope X", ROUND(h.slope_y,4) "Slope Y",
                  ROUND(h.hyst_pos0_x,3) "Hyst X", ROUND(h.hyst_pos0_y,3) "Hyst Y"
           FROM head_health h JOIN machines m ON m.id=h.machine_id ORDER BY h.ts DESC LIMIT 60""")]
    return {"title": "Maintenance Analysis", "subtitle": f"{start} to {end}",
            "sections": [_sec("Maintenance risk", rows),
                         _sec("Recurring alarms", al["top"][:20]),
                         _sec("Scan-head calibration history", hcm)]}


def alert_history_report(days=30):
    start, end = _rng(days)
    rows = [dict(r) for r in db.q(
        """SELECT a.started_at "Started", m.name "Machine", a.severity "Severity", a.type "Type",
                  a.title "Title", a.status "Status", a.ack_by "Acknowledged by", a.ack_at "Ack at",
                  a.resolved_at "Resolved at", ROUND(a.duration_seconds/60.0,1) "Duration min",
                  a.resolution "Resolution"
           FROM alerts a LEFT JOIN machines m ON m.id=a.machine_id
           WHERE a.started_at >= ? ORDER BY a.started_at DESC""", (start,))]
    s = alerts.summary(days=days)
    return {"title": "Alert Report", "subtitle": f"{start} to {end}",
            "sections": [_sec("By type", s["by_type"]), _sec("By machine", s["by_machine"]),
                         _sec("History", rows)]}


def connectivity_report(days=7):
    start, end = _rng(days)
    rows = [dict(r) for r in db.q(
        """SELECT m.name "Machine", m.host "Host", m.connection_method "Method",
                  COUNT(*) "Checks", SUM(c.ok) "OK", ROUND(AVG(c.ok)*100,1) "Uptime %",
                  ROUND(AVG(c.latency_ms),0) "Avg latency ms", MAX(c.ts) "Last check"
           FROM connection_log c JOIN machines m ON m.id=c.machine_id
           WHERE c.ts >= ? AND c.kind='heartbeat' GROUP BY c.machine_id ORDER BY "Uptime %" ASC""",
        (start,))]
    fails = [dict(r) for r in db.q(
        """SELECT c.ts "Time", m.name "Machine", c.kind "Kind", c.detail "Detail"
           FROM connection_log c JOIN machines m ON m.id=c.machine_id
           WHERE c.ts >= ? AND c.ok=0 ORDER BY c.ts DESC LIMIT 200""", (start,))]
    return {"title": "Connectivity & Data Ingestion Report", "subtitle": f"{start} to {end}",
            "sections": [_sec("Machine connectivity", rows), _sec("Recent failures", fails)]}


def optics_report():
    """Every laser-recipe deviation from the fleet standard, worst first."""
    d = analytics.optics_drift()
    summary = [{"Machine": b["machine"], "Model": b["model"], "Deviations": b["total"],
                "Critical": b["critical"], "Units/hour": b["units_per_hour"],
                "Model median": b["model_median_uph"],
                "vs model %": b["vs_model_pct"],
                "Worst speed deviation %": b["speed_loss_pct"]} for b in d["by_machine"]]
    detail = [{"Machine": f["machine"], "Model": f["model"], "Preset": f["preset"],
               "Parameter": f["param"], "This machine": f["value"], "Fleet standard": f["standard"],
               "Peers agreeing": f"{f['agree']}/{f['peers']}", "Delta %": f["delta_pct"],
               "Area": f["area"], "Severity": f["severity"], "Note": f["note"]}
              for f in d["findings"]]
    cov = [{"Metric": "Machines with optics captured", "Value": d["machines_with_optics"]},
           {"Metric": "Machines showing drift", "Value": d["machines_with_drift"]},
           {"Metric": "Presets compared", "Value": d["presets"]},
           {"Metric": "Parameters compared", "Value": d["parameters"]},
           {"Metric": "Total deviations", "Value": len(d["findings"])}]
    return {"title": "Laser Optics Configuration Drift", "subtitle": dt.date.today().isoformat(),
            "sections": [_sec("Coverage", cov), _sec("By machine", summary),
                         _sec("All deviations", detail)]}


REPORTS = {
    "optics": ("Laser Optics Configuration Drift", optics_report),
    "daily_production": ("Daily Laser Production Report", daily_production),
    "shift": ("Shift Performance Report", shift_report),
    "utilization": ("Machine Utilization Report", utilization_report),
    "downtime": ("Downtime Report", downtime_report),
    "alarms": ("Alarm Report", alarm_report),
    "health": ("Machine Health Report", health_report),
    "weekly": ("Weekly Performance Report", weekly_report),
    "monthly": ("Monthly Management Report", monthly_report),
    "loss": ("Production Loss Analysis", loss_report),
    "capacity": ("Capacity Analysis", capacity_report),
    "maintenance": ("Maintenance Analysis", maintenance_report),
    "alert_history": ("Alert Report", alert_history_report),
    "connectivity": ("Connectivity & Data Ingestion Report", connectivity_report),
}


def build(name, **kw):
    if name not in REPORTS:
        raise KeyError(name)
    fn = REPORTS[name][1]
    import inspect
    accepted = set(inspect.signature(fn).parameters)
    return fn(**{k: v for k, v in kw.items() if k in accepted and v is not None})


def _r(v):
    return round(v, 1) if isinstance(v, float) else v


# ---------------------------------------------------------------- renderers

HEAD_FILL = PatternFill("solid", fgColor="0F2A44")
TITLE_FONT = Font(size=14, bold=True, color="0F2A44")


def to_xlsx(report) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for s in report["sections"]:
        ws = wb.create_sheet((s["name"] or "Sheet")[:31])
        ws.append([report["title"]])
        ws["A1"].font = TITLE_FONT
        ws.append([f"{s['name']} — {report.get('subtitle', '')}"])
        ws.append([])
        ws.append([str(c) for c in s["columns"]])
        hdr = ws.max_row
        for cell in ws[hdr]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(horizontal="center")
        for row in s["rows"]:
            ws.append(["" if v is None else v for v in row])
        ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
        for i, c in enumerate(s["columns"], 1):
            width = max(len(str(c)), *(len(str(r[i - 1])) for r in s["rows"])) if s["rows"] else len(str(c))
            ws.column_dimensions[get_column_letter(i)].width = min(42, max(10, width + 2))
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=report["title"])
    ss = getSampleStyleSheet()
    story = [Paragraph(f"<b>{report['title']}</b>", ss["Title"]),
             Paragraph(f"{BRAND} — {report.get('subtitle', '')} — generated {db.now()}", ss["Normal"]),
             Spacer(1, 6 * mm)]
    for s in report["sections"]:
        story.append(Paragraph(f"<b>{s['name']}</b>", ss["Heading2"]))
        if not s["rows"]:
            story.append(Paragraph("No data for this period.", ss["Normal"]))
            story.append(Spacer(1, 4 * mm))
            continue
        data = [[Paragraph(f"<b>{c}</b>", ss["BodyText"]) for c in s["columns"]]]
        LIMIT = 2000                      # keep the PDF openable, but never drop rows silently
        for row in s["rows"][:LIMIT]:
            data.append([Paragraph(_txt(v), ss["BodyText"]) for v in row])
        if len(s["rows"]) > LIMIT:
            note = [""] * len(s["columns"])
            note[0] = f"... {len(s['rows']) - LIMIT} more rows not shown - export to Excel for the full set"
            data.append([Paragraph(f"<i>{_txt(v)}</i>", ss["BodyText"]) for v in note])
        t = Table(data, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F2A44")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C4D0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F8")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(t)
        story.append(Spacer(1, 6 * mm))
    doc.build(story)
    return buf.getvalue()


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float):
        v = round(v, 2)
    return (str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))[:200]
